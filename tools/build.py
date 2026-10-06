#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build.py — مُولِّد أوراق العمل وأدلة المدرّب
================================================
المصادر:  source/*.md        (بصيغة DSL بسيطة)
المخرجات: worksheets/*.html + *.pdf     أوراق الطلاب
          solutions/*.html  + *.pdf     نماذج الحلول
          trainer/*.html    + *.pdf     أدلة المدرّب
          docx/*.docx                    نسخ قابلة للتعديل

الاستخدام:
    python3 tools/build.py ws1            # ملف واحد
    python3 tools/build.py                # الكل
    python3 tools/build.py --no-pdf       # HTML + DOCX بدون PDF
    python3 tools/build.py --no-docx      # بدون DOCX (أسرع)
"""

from __future__ import annotations

import html
import json
import os
import re
import subprocess
import sys
import textwrap
from pathlib import Path

import markdown

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "source"
OUT = {
    "worksheet": ROOT / "worksheets",
    "solution": ROOT / "solutions",
    "trainer": ROOT / "trainer",
}
WORK = ROOT / "_work" / "build"
WORK.mkdir(parents=True, exist_ok=True)

MD_EXTENSIONS = ["tables", "attr_list", "md_in_html", "sane_lists", "def_list"]

# ---------------------------------------------------------------------------
# حماية كتل HTML من ماركداون
# ---------------------------------------------------------------------------

_registry: dict[str, str] = {}
_counter = [0]


def store(html_str: str) -> str:
    _counter[0] += 1
    token = f"<!--H{_counter[0]}-->"
    _registry[token] = html_str
    return token


def restore(text: str) -> str:
    for token, html_str in _registry.items():
        if token in text:
            text = text.replace(token, html_str)
    return text


def esc(t) -> str:
    return html.escape(str(t), quote=True)


FENCE_RE = re.compile(r"```(?:[a-zA-Z0-9_-]*)\n(.*?)```", re.S)


def md(text: str) -> str:
    """ماركداون + دعم الكتل المحدودة بـ ``` كصناديق إدخال."""
    text = textwrap.dedent(text).strip()

    def repl(m):
        body = m.group(1).rstrip()
        inner = f'<span class="ph">{html.escape(body)}</span>' if body.strip() else ""
        return "\n\n" + store(f'<pre class="plain">{inner}\n</pre>') + "\n\n"

    text = FENCE_RE.sub(repl, text)
    return markdown.markdown(text, extensions=MD_EXTENSIONS)


def md_restore(text: str) -> str:
    return restore(md(text))


# ---------------------------------------------------------------------------
# الكتل
# ---------------------------------------------------------------------------

OPEN_RE = re.compile(r"^:::(?P<name>[a-zA-Z0-9_-]+)(?:\{(?P<attrs>.*)\})?\s*$")
CLOSE_RE = re.compile(r"^:::\s*$")


def parse_attrs(raw: str | None) -> dict:
    if not raw:
        return {}
    attrs = {}
    for m in re.finditer(r'([a-zA-Z0-9_-]+)(?:\s*=\s*(?:"([^"]*)"|\'([^\']*)\'|([^\s]+)))?', raw):
        val = next((g for g in m.groups()[1:] if g is not None), "1")
        attrs[m.group(1)] = val
    return attrs


def tools_row(chips: str) -> str:
    return f'<div class="task-tools">{chips}</div>' if chips else ''


def b_task(a: dict, inner: str) -> str:
    chips = ""
    if a.get("time"):
        chips += f'<span class="chip k">⏱ {esc(a["time"])}</span>'
    if a.get("level"):
        chips += f'<span class="chip o">◆ {esc(a["level"])}</span>'
    if a.get("pts"):
        chips += f'<span class="chip t">★ {esc(a["pts"])} نقطة</span>'
    for t in [x for x in a.get("tools", "").split("|") if x.strip()]:
        chips += f'<span class="chip">{esc(t.strip())}</span>'
    return (
        '<div class="task"><div class="task-head">'
        f'<div class="task-num">{esc(a.get("num", ""))}</div>'
        f'<div class="task-titles"><div class="ar">{esc(a.get("ar", ""))}</div>'
        f'<div class="en">{esc(a.get("en", ""))}</div></div></div>'
        f'<div class="task-body">{tools_row(chips)}{md_restore(inner)}</div></div>'
    )


def _hl(inp: str) -> str:
    out = []
    for line in inp.split("\n"):
        e = html.escape(line)
        if line.startswith("$ "):
            out.append(f'<span class="p">$</span> <span class="k">{html.escape(line[2:])}</span>')
        elif line.startswith("#") or line.startswith("// "):
            out.append(f'<span class="c">{e}</span>')
        elif re.match(r"^(fatal:|error:|CONFLICT|Auto-merging|remote: error:)", line):
            out.append(f'<span class="s">{e}</span>')
        elif (re.match(r"^(remote:|hint:|warning:|note:)", line)
              or line.startswith("<<<<<<<") or line.startswith("=======") or line.startswith(">>>>>>>")):
            out.append(f'<span class="c">{e}</span>')
        elif line.startswith(("To ", "From ", " ! ", "branch ", " * ", "Everything ")):
            out.append(f'<span class="o">{e}</span>')
        else:
            out.append(f'<span class="o">{e}</span>')
    return "\n".join(out)


def b_term(a: dict, inner: str) -> str:
    return (
        f'<div class="term"><div class="bar"><i class="r"></i><i class="y"></i><i class="g"></i>'
        f'<span class="t">{esc(a.get("title", "Terminal"))}</span></div><pre>{_hl(inner.rstrip())}</pre></div>'
    )


def b_code(a: dict, inner: str) -> str:
    return f'<pre class="plain">{html.escape(inner.strip())}</pre>'


def b_box(a: dict, inner: str) -> str:
    kind = esc(a.get("type", "info"))
    ar, en = esc(a.get("ar", "")), esc(a.get("en", ""))
    head = f'<div class="h">{ar}<span class="en">{en}</span></div>' if (ar or en) else ""
    return f'<div class="box {kind}">{head}{md_restore(inner)}</div>'


def b_bi(a: dict, inner: str) -> str:
    tone = f'tone-{a["tone"]}' if a.get("tone") else ""
    ar_part, en_part = inner.split("---EN---", 1) if "---EN---" in inner else (inner, "")
    en_html = f'<div class="eng">{md_restore(en_part)}</div>' if en_part.strip() else ""
    return f'<div class="bi {tone}"><div class="ar">{md_restore(ar_part)}</div>{en_html}</div>'


def b_write(a: dict, inner: str) -> str:
    n = int(a.get("n", "3"))
    label = esc(a.get("label", ""))
    lab = f'<div class="lbl">{label}</div>' if label else ""
    line_div = '<div class="ln"></div>'
    return f'<div class="write">{lab}{line_div * n}</div>'


def b_goal(a: dict, inner: str) -> str:
    return (
        f'<div class="goal-list"><strong>{esc(a.get("ar", "أهداف الجلسة"))}'
        f'<span class="muted tiny"> / {esc(a.get("en", "Objectives"))}</span></strong>{md_restore(inner)}</div>'
    )


def b_evidence(a: dict, inner: str) -> str:
    return (
        f'<div class="evidence"><div class="h">{esc(a.get("ar", "دليل التسليم المطلوب"))}'
        f'<span class="en">/ {esc(a.get("en", "Required evidence"))}</span></div>{md_restore(inner)}</div>'
    )


def b_check(a: dict, inner: str) -> str:
    body = md_restore(inner).replace("<ul>", '<ul class="checklist">', 1)
    head = (f'<strong>{esc(a.get("ar", "قائمة الفحص قبل التسليم"))}'
            f'<span class="muted tiny"> / {esc(a.get("en", "Checklist"))}</span></strong>')
    return f'<div class="write">{head}{body}</div>'


def b_scale(a: dict, inner: str) -> str:
    items = a.get("items", "لم أبدأ|بدأت|أنجزت|أتقنت").split("|")
    cells = "".join(f'<div class="s"><span class="b">{i+1}</span>{esc(it.strip())}</div>' for i, it in enumerate(items))
    label = esc(a.get("label", "قيّم نفسك / Self-assessment"))
    return f'<div class="write"><div class="lbl">{label}</div><div class="scale">{cells}</div></div>'


def b_badge(a: dict, inner: str) -> str:
    return f'<p class="center"><span class="stamp {esc(a.get("type",""))}">{esc(a.get("ar",""))}</span></p>'


def b_band(a: dict, inner: str) -> str:
    return (f'<div class="band {esc(a.get("type",""))}"><span class="ar">{esc(a.get("ar",""))}</span>'
            f'<span class="en">{esc(a.get("en",""))}</span></div>')


BLOCKS = {
    "task": b_task, "term": b_term, "code": b_code, "box": b_box, "bi": b_bi,
    "write": b_write, "lines": b_write, "goal": b_goal, "evidence": b_evidence,
    "check": b_check, "scale": b_scale, "badge": b_badge, "band": b_band,
}
RAW_BLOCKS = {"term", "code"}


SECTION_BLOCKS = {"band", "badge"}


def _render_node(node: dict) -> str:
    """يحوّل عقدة كتلة واحدة إلى HTML (مع معالجة الاستثناءات)."""
    fn = BLOCKS.get(node["name"])
    inner = "\n".join(node["buf"])
    if not fn:
        return md(inner)
    if node["name"] in RAW_BLOCKS:
        return fn(node["attrs"], inner)
    return fn(node["attrs"], render_blocks(inner))


def _close_top(stack: list, out: list) -> None:
    """يغلق أعلى كتلة في المكدس ويضع ناتجها في العنصر الأب."""
    node = stack.pop()
    rendered = _render_node(node)
    (stack[-1]["buf"] if stack else out).append(store(rendered))


def render_blocks(text: str) -> str:
    """يحوّل الكتل المتشعّبة إلى HTML.

    المتسامح: يُغلق الكتلة السابقة تلقائياً عند بدء كتلة من النوع نفسه،
    أو عند بدء كتلة قسم (band/badge)، أو عند نهاية الملف.
    """
    out: list[str] = []
    stack: list[dict] = []
    for line in text.split("\n"):
        m = OPEN_RE.match(line)
        if m and not CLOSE_RE.match(line):
            name = m.group("name")
            if name in SECTION_BLOCKS:
                while stack:
                    _close_top(stack, out)
            elif stack and stack[-1]["name"] == name:
                _close_top(stack, out)
            stack.append({"name": name, "attrs": parse_attrs(m.group("attrs")), "buf": []})
            continue
        if CLOSE_RE.match(line) and stack:
            _close_top(stack, out)
            continue
        (stack[-1]["buf"] if stack else out).append(line)
    while stack:
        _close_top(stack, out)
    return "\n".join(out)


# ---------------------------------------------------------------------------
# القالب
# ---------------------------------------------------------------------------

DOC_TEMPLATE = """<!DOCTYPE html>
<html lang="ar" dir="rtl">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<link rel="stylesheet" href="{css}">
</head>
<body>
<div class="sheet">
{dochead}
{body}
{footer}
</div>
</body>
</html>
"""

IDENTITY = """
<div class="identity">
  <div class="fld"><span class="k">اسم الطالب / Student</span><span class="dots"></span></div>
  <div class="fld"><span class="k">اسم المستخدم / GitHub username</span><span class="dots"></span></div>
  <div class="fld"><span class="k">المجموعة / Group</span><span class="dots"></span></div>
  <div class="fld"><span class="k">التاريخ / Date</span><span class="dots"></span></div>
</div>
"""


def doc_head(meta: dict, kind: str) -> str:
    cells = "".join(f'<div><span class="k">{esc(k)}</span><span class="v">{esc(v)}</span></div>'
                    for k, v in meta.get("meta", []))
    tone = ' style="background:#7A1220"' if kind == "solution" else (
        ' style="background:#0E5E4A"' if kind == "trainer" else "")
    return (
        '<div class="doc-head"><div class="top"' + tone + '>'
        f'<div class="badge"><div class="n">{esc(meta.get("badge_n",""))}</div>'
        f'<div class="l">{esc(meta.get("badge_l","ورقة عمل"))}</div></div>'
        f'<div class="titles"><div class="kicker">{esc(meta.get("kicker",""))}</div>'
        f'<h1>{meta.get("title","")}</h1>'
        f'<div class="sub">{esc(meta.get("en_title",""))}</div></div></div>'
        f'<div class="meta">{cells}</div></div>'
    )


def parse_source(path: Path) -> tuple[dict, str]:
    text = path.read_text(encoding="utf-8")
    meta: dict = {}
    if text.lstrip().startswith("---"):
        stripped = text.lstrip()
        end = stripped.find("\n---", 3)
        if end != -1:
            meta = json.loads(stripped[3:end])
            text = stripped[end + 4:]
    return meta, text


def build_html(src_file: Path, kind: str) -> Path:
    _registry.clear()
    _counter[0] = 0
    meta, body = parse_source(src_file)
    body_html = restore(md(render_blocks(body)))
    if kind == "worksheet" and meta.get("identity", True):
        body_html = IDENTITY + body_html
    banner = {
        "solution": '<div class="box err"><div class="h">نسخة المدرّب — نماذج الحلول والتقييم '
                    '<span class="en">/ Instructor copy — do not distribute</span></div></div>',
        "trainer": '<div class="box ok"><div class="h">دليل المدرّب '
                   '<span class="en">/ Trainer guide</span></div></div>',
    }.get(kind, "")
    footer = ""
    if meta.get("footer"):
        footer = (f'<div class="footer-note"><span>{esc(meta["footer"])}</span>'
                  f'<span class="en">{esc(meta.get("en_footer",""))}</span></div>')
    doc = DOC_TEMPLATE.format(
        title=esc(meta.get("title", src_file.stem)),
        css="../assets/print.css",
        dochead=doc_head(meta, kind),
        body=banner + body_html,
        footer=footer,
    )
    target_dir = OUT[kind]
    target_dir.mkdir(parents=True, exist_ok=True)
    out = target_dir / (src_file.stem + ".html")
    out.write_text(doc, encoding="utf-8")
    return out


# ---------------------------------------------------------------------------
# PDF / DOCX
# ---------------------------------------------------------------------------

PDF_SCRIPT = r"""
import chromium from '@sparticuz/chromium';
import puppeteer from 'puppeteer-core';
const files = process.argv.slice(2);
const browser = await puppeteer.launch({
  args: [...chromium.args, '--no-sandbox', '--disable-setuid-sandbox', '--font-render-hinting=none'],
  executablePath: await chromium.executablePath(),
  headless: true,
});
for (const f of files) {
  const pdf = f.replace(/\.html$/, '.pdf');
  const page = await browser.newPage();
  await page.goto('file://' + f, { waitUntil: 'networkidle0' });
  await page.evaluate(() => document.fonts.ready);
  await new Promise(r => setTimeout(r, 200));
  await page.pdf({
    path: pdf, format: 'A4', printBackground: true,
    margin: { top: '11mm', bottom: '13mm', left: '10mm', right: '10mm' },
    displayHeaderFooter: true,
    headerTemplate: '<div></div>',
    footerTemplate: '<div style="width:100%;font-size:7pt;color:#8494A8;font-family:Arial,sans-serif;padding:0 10mm;text-align:center;">Git &amp; GitHub Workshop Workbook &nbsp;·&nbsp; Page <span class="pageNumber"></span> / <span class="totalPages"></span></div>'
  });
  await page.close();
  console.log('PDF v', pdf);
}
await browser.close();
"""


def run_pdfs(html_files: list[Path]) -> None:
    if not html_files:
        return
    script = WORK / "pdf.mjs"
    script.write_text(PDF_SCRIPT, encoding="utf-8")
    if not (WORK / "node_modules").exists():
        print("…… تثبيت chromium/puppeteer مرة واحدة")
        subprocess.run(["npm", "install", "--silent", "@sparticuz/chromium", "puppeteer-core"],
                       cwd=WORK, check=True)
    env = os.environ.copy()
    env["LD_LIBRARY_PATH"] = "/tmp/chlibs/lib:" + env.get("LD_LIBRARY_PATH", "")
    subprocess.run(["node", str(script), *[str(f) for f in html_files]], cwd=WORK, env=env, check=True)


def rasterize_assets() -> None:
    """يحوّل رسومات SVG إلى PNG لاستخدامها في DOCX (مرة واحدة)."""
    png_dir = ROOT / "assets" / "img" / "png"
    svgs = sorted((ROOT / "assets" / "img").glob("*.svg"))
    fresh = [s for s in svgs if not (png_dir / (s.stem + ".png")).exists()
             or (png_dir / (s.stem + ".png")).stat().st_mtime < s.stat().st_mtime]
    if not fresh:
        png_dir.mkdir(parents=True, exist_ok=True)
        return
    script = WORK / "rasterize.mjs"
    script.write_text((ROOT / "tools" / "rasterize.mjs").read_text(encoding="utf-8"), encoding="utf-8")
    if not (WORK / "node_modules").exists():
        subprocess.run(["npm", "install", "--silent", "@sparticuz/chromium", "puppeteer-core"],
                       cwd=WORK, check=True)
    env = os.environ.copy()
    env["LD_LIBRARY_PATH"] = "/tmp/chlibs/lib:" + env.get("LD_LIBRARY_PATH", "")
    subprocess.run(["node", str(script), str(png_dir), *[str(s) for s in fresh]],
                   cwd=WORK, env=env, check=True)


def run_docx(html_file: Path) -> None:
    try:
        import pypandoc
    except ImportError:
        return
    dest_dir = ROOT / "docx"
    dest_dir.mkdir(exist_ok=True)
    dest = dest_dir / (html_file.stem + ".docx")
    pandoc = pypandoc.get_pandoc_path()
    env = os.environ.copy()
    env["PATH"] = str(Path(pandoc).parent) + os.pathsep + env.get("PATH", "")
    # نسخة وسيطة تشير إلى صور PNG مطلقة حتى يستطيع pandoc تضمينها
    rasterize_assets()
    src = html_file.read_text(encoding="utf-8")
    src = re.sub(r'\.\./assets/img/([a-zA-Z0-9_-]+)\.svg',
                 lambda m: (ROOT / "assets" / "img" / "png" / (m.group(1) + ".png")).as_uri(), src)
    tmp = WORK / (html_file.stem + ".docx.html")
    tmp.write_text(src, encoding="utf-8")
    subprocess.run([pandoc, "-f", "html", "-t", "docx", "--standalone", "-o", str(dest), str(tmp)],
                   check=True, env=env)
    fix_docx_rtl(dest)


def fix_docx_rtl(path: Path) -> None:
    try:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        from docx.shared import Pt
    except ImportError:
        return
    doc = Document(str(path))

    def rtl_para(p):
        pPr = p._p.get_or_add_pPr()
        if pPr.find(qn("w:bidi")) is None:
            pPr.append(OxmlElement("w:bidi"))
        for r in p.runs:
            rPr = r._r.get_or_add_rPr()
            if rPr.find(qn("w:rtl")) is None:
                rPr.append(OxmlElement("w:rtl"))
            if r.font.size is None:
                r.font.size = Pt(11)

    for p in doc.paragraphs:
        rtl_para(p)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    rtl_para(p)
        tblPr = t._tbl.tblPr
        if tblPr.find(qn("w:bidiVisual")) is None:
            tblPr.append(OxmlElement("w:bidiVisual"))
    try:
        doc.styles["Normal"].font.name = "Arial"
    except Exception:
        pass
    doc.save(str(path))


# ---------------------------------------------------------------------------
def main() -> None:
    args = sys.argv[1:]
    no_pdf = "--no-pdf" in args
    no_docx = "--no-docx" in args
    patterns = [a for a in args if not a.startswith("--")]
    files = sorted(SRC.glob("*.md"))
    if patterns:
        files = [f for f in files if any(p in f.name for p in patterns)]
    if not files:
        print("لا ملفات مصدر مطابقة.")
        return

    html_files = []
    for f in files:
        kind = "solution" if "solution" in f.name else ("trainer" if "trainer" in f.name else "worksheet")
        h = build_html(f, kind)
        html_files.append(h)
        print(f"HTML ✓ {h.relative_to(ROOT)}")
        if not no_docx:
            try:
                run_docx(h)
                print(f"DOCX ✓ docx/{h.stem}.docx")
            except Exception as e:  # noqa: BLE001
                print("  ! تعذّر DOCX:", e)
    if not no_pdf:
        run_pdfs(html_files)
    print("\nاكتمل البناء.")


if __name__ == "__main__":
    main()
